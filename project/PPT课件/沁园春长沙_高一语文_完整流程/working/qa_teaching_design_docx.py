# -*- coding: utf-8 -*-
"""教学设计 DOCX 格式 QA：字体、字号、颜色、Markdown 残留。"""
import re
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

path = r"E:\codex\jiaosdhi\eyes\project\PPT课件\赤壁赋_高一语文_完整流程\outputs\teaching_design.docx"
doc = Document(path)

errors = []

# 1. 字体/字号/颜色校验
def check_para(p, where):
    for r in p.runs:
        if not r.text.strip():
            continue
        name = r.font.name
        ea = r._element.rPr.rFonts.get(qn("w:eastAsia")) if r._element.rPr is not None and r._element.rPr.rFonts is not None else None
        size = r.font.size
        color = r.font.color.rgb if r.font.color and r.font.color.rgb else None
        if name != "宋体":
            errors.append(f"{where} 字体={name!r} 文本={r.text[:20]!r}")
        if ea != "宋体":
            errors.append(f"{where} 中文字体={ea!r} 文本={r.text[:20]!r}")
        if size is None or size.pt != 10.5:
            errors.append(f"{where} 字号={size.pt if size else None} 文本={r.text[:20]!r}")
        if color is None or str(color) != "000000":
            errors.append(f"{where} 颜色={color} 文本={r.text[:20]!r}")

for i, p in enumerate(doc.paragraphs):
    check_para(p, f"段落{i}")

for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                check_para(p, f"表{ti}行{ri}列{ci}")

# 2. Markdown 残留扫描
md_patterns = [
    (r"^#{1,6}\s", "ATX标题"),
    (r"^\s*[-*+]\s+", "列表符号"),
    (r"\*\*", "加粗星号"),
    (r"`", "反引号"),
    (r"^\s*>\s", "引用符号"),
    (r"^\s*\|.*\|", "Markdown表格"),
    (r"^\s*-{3,}\s*$", "分隔线"),
    (r"!\[.*?\]\(", "图片语法"),
    (r"\[.*?\]\(.*?\)", "链接语法"),
]
texts = [p.text for p in doc.paragraphs]
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            texts.extend(p.text for p in cell.paragraphs)

md_hits = []
for txt in texts:
    for pat, label in md_patterns:
        if re.search(pat, txt, re.M):
            md_hits.append(f"{label}: {txt[:50]!r}")

# 3. 统计
print("=== 段落数 ===", len(doc.paragraphs))
print("=== 表格数 ===", len(doc.tables))
print("=== 格式错误数 ===", len(errors))
for e in errors[:15]:
    print("  !", e)
print("=== Markdown 残留数 ===", len(md_hits))
for m in md_hits[:10]:
    print("  !", m)

# 4. 结构完整性
required = ["一、基本信息", "二、教材与学情分析", "三、教学目标", "四、教学重点与难点",
            "五、教学方法与资源", "六、教学过程", "八、板书设计", "九、作业设计",
            "十、AI赋能设计", "十一、教学评价"]
all_text = "\n".join(texts)
missing = [s for s in required if s not in all_text]
print("=== 缺失章节 ===", missing if missing else "无")

ok = not errors and not md_hits and not missing
print("=== QA 结论 ===", "PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
