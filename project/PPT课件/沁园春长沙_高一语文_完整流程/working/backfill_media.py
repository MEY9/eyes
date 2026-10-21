# -*- coding: utf-8 -*-
"""可编辑课件的媒体回填后处理。

在 `editppt run finalize` 之后运行，做两件 finalize 的合成器不负责的事：

1. 把 page_002 的视频槽位图片替换为真正嵌入的 MP4（同位置同大小，带 poster）。
2. 把 page_024 底部的静态说明文字变成指向 HTML 互动课件的超链接入口。

这样静态备用版天然保留（poster 图 + 六只条框），交互版在支持的环境下点击即可打开。

用法：
    python working/backfill_media.py
    python working/backfill_media.py --dry-run
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError as exc:  # pragma: no cover
    raise SystemExit(f"python-pptx is required: {exc}")


ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "outputs" / "赤壁赋_可编辑版.pptx"
VIDEO = ROOT / "resources" / "ai_video" / "ai-video-01-1080p.mp4"
POSTER = ROOT / "resources" / "ai_video" / "ai-video-01-poster.png"
HTML = ROOT / "html_embeds" / "ai-html-01" / "ai-html-01.html"
REPORT = ROOT / "qa" / "editable_deck" / "media_backfill_report.json"

VIDEO_SLOT_ID = "AI-VIDEO-01_VIDEO_SLOT"
INTERACTIVE_HINT = "互动"


def find_video_slot(slide):
    """Locate the page_002 video placeholder picture.

    The placeholder is authored either with the English id
    ``AI-VIDEO-01_VIDEO_SLOT`` or with its Chinese label 「AI 视频槽位」,
    depending on which run produced it, so both spellings are accepted.
    """
    markers = (VIDEO_SLOT_ID, "AI 视频槽位", "视频槽位")
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        name = shape.name or ""
        if any(marker in name for marker in markers):
            return shape
    return None


def find_interactive_entry(slide):
    """Locate the page_024 footer note that acts as the interactive entry."""
    candidates = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text or ""
        if INTERACTIVE_HINT in text:
            candidates.append(shape)
    if not candidates:
        return None
    # Pick the bottom-most note line when several mention the interaction.
    return max(candidates, key=lambda s: (s.top or 0))


def backfill_video(slide, video: Path, poster: Path, dry_run: bool = False):
    """Replace the placeholder picture with an embedded MP4 at the same box."""
    slot = find_video_slot(slide)
    if slot is None:
        return {"status": "skipped", "reason": "video slot picture not found"}

    left, top, width, height = slot.left, slot.top, slot.width, slot.height
    sp_tree = slide.shapes._spTree
    slot_index = list(sp_tree).index(slot._element)

    if dry_run:
        return {
            "status": "planned",
            "target_shape": slot.name,
            "box_pt": [round(v / 12700, 2) for v in (left, top, width, height)],
            "video": str(video),
            "poster": str(poster),
            "insert_index": slot_index,
        }

    sp_tree.remove(slot._element)
    movie = slide.shapes.add_movie(
        str(video),
        left,
        top,
        width,
        height,
        poster_frame_image=str(poster),
        mime_type="video/mp4",
    )
    # Keep the original stacking order: the video stays behind the text layer.
    movie_el = movie._element
    sp_tree.remove(movie_el)
    sp_tree.insert(slot_index, movie_el)

    return {
        "status": "done",
        "replaced_shape": slot.name,
        "box_pt": [round(v / 12700, 2) for v in (left, top, width, height)],
        "video": str(video),
        "poster": str(poster),
        "z_index_position": slot_index,
    }


def backfill_html_link(slide, html: Path, dry_run: bool = False):
    """Turn the static footer note into a clickable entry to the HTML module."""
    entry = find_interactive_entry(slide)
    if entry is None:
        return {"status": "skipped", "reason": "interactive footer note not found"}

    frame = entry.text_frame
    original = frame.text or ""
    target = html.name
    if html.parent != ROOT:
        target = str(html.relative_to(ROOT)).replace("\\", "/")

    if dry_run:
        return {"status": "planned", "current_text": original, "link_target": target}

    # Rewrite the note as an explicit entry line, then attach the hyperlink.
    new_text = "点击此处打开互动课件（主客问答 · 配对）；环境不支持时本页静态备用内容可直接放映。"
    paragraphs = frame.paragraphs
    if paragraphs:
        first = paragraphs[0]
        for run in list(first.runs):
            run._r.getparent().remove(run._r)
        if first.runs:
            first.runs[0].text = new_text
        else:
            run = first.add_run()
            run.text = new_text
        run = first.runs[0]
        run.hyperlink.address = target
        # Drop the remaining paragraphs so the note stays a single line.
        for para in paragraphs[1:]:
            para._p.getparent().remove(para._p)
    else:
        run = frame.paragraphs[0].add_run()
        run.text = new_text
        run.hyperlink.address = target

    return {
        "status": "done",
        "previous_text": original,
        "new_text": new_text,
        "link_target": target,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Backfill media into the finalized editable deck.")
    parser.add_argument("--dry-run", action="store_true", help="Report planned changes without writing.")
    parser.add_argument("--no-backup", action="store_true", help="Skip creating a timestamped backup.")
    parser.add_argument(
        "--out",
        help="Output path. Defaults to overwriting the deck in place; use a different "
             "path when the deck is still open in PowerPoint.",
    )
    args = parser.parse_args()
    out_path = Path(args.out) if args.out else DECK

    missing = [path for path in (DECK, VIDEO, POSTER, HTML) if not path.exists()]
    if missing:
        for path in missing:
            print(f"Error: missing required file: {path}", file=sys.stderr)
        return 1

    report = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "deck": str(DECK),
        "dry_run": args.dry_run,
        "actions": [],
    }

    backup_path = None
    if not args.dry_run and not args.no_backup:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = DECK.with_name(f"{DECK.stem}.bak_{stamp}{DECK.suffix}")
        shutil.copy2(DECK, backup_path)
        report["backup"] = str(backup_path)

    report["output"] = str(out_path)

    prs = Presentation(str(DECK))

    # --- Slide 2: swap the placeholder picture for an embedded MP4 ---------
    slide2 = prs.slides[1]
    video_result = backfill_video(slide2, VIDEO, POSTER, dry_run=args.dry_run)
    report["actions"].append({"page": 2, "kind": "embedded-video", **video_result})

    # --- Slide 24: make the footer note an entry to the HTML module ---------
    slide24 = prs.slides[23]
    html_result = backfill_html_link(slide24, HTML, dry_run=args.dry_run)
    report["actions"].append({"page": 24, "kind": "html-interactive-entry", **html_result})

    if not args.dry_run:
        try:
            prs.save(str(out_path))
            report["written"] = True
        except PermissionError:
            # The deck is open in PowerPoint. Fall back to a sibling copy so the
            # work is never lost, and report the blocked in-place write.
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            fallback = out_path.with_name(f"{out_path.stem}_媒体回填_{stamp}{out_path.suffix}")
            prs.save(str(fallback))
            report["written"] = True
            report["output"] = str(fallback)
            report["blocked_in_place"] = str(out_path)
            report["blocked_reason"] = "file is open in another program (PowerPoint/预览器)"

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
