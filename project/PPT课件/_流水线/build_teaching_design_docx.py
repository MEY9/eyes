# -*- coding: utf-8 -*-
"""通用教学设计 DOCX 生成器（语文课件流水线）。

从 working/lesson_packet.json 读取结构化数据，输出正式教学设计 Word：
黑色、五号、宋体、普通 Word 表格，不使用 Markdown 装饰符号。

用法：
    python _流水线/build_teaching_design_docx.py \
        --lesson-packet <项目>/working/lesson_packet.json \
        --out <项目>/outputs/teaching_design.docx
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "宋体"
SIZE = Pt(10.5)  # 五号
BLACK = RGBColor(0, 0, 0)


def _run(p, text, bold=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = SIZE
    r.font.bold = bold
    r.font.color.rgb = BLACK
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return r


class Doc:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(3.18)
        sec.right_margin = Cm(3.18)
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = SIZE
        normal.font.color.rgb = BLACK
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(0)

    def title(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        _run(p, text, bold=True)

    def h1(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        _run(p, text, bold=True)

    def h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Pt(21)
        _run(p, text, bold=True)

    def body(self, text, indent=True):
        p = self.doc.add_paragraph()
        if indent:
            p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.space_after = Pt(3)
        _run(p, text)

    def table(self, headers, rows, widths=None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = ""
            p = hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, h, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                _run(p, str(val))
        if widths:
            for r in t.rows:
                for i, w in enumerate(widths):
                    r.cells[i].width = Cm(w)
        return t

    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)


def build(packet: dict, d: Doc):
    nv = packet.get("narrative", {})
    title = packet.get("title", "")
    d.title(f"《{title}》教学设计")

    # 一、基本信息
    d.h1("一、基本信息")
    d.table(
        ["项目", "内容"],
        [
            ["课题", title],
            ["作者", packet.get("author", "—")],
            ["学段学科", f'{packet.get("grade", "")}{packet.get("subject", "")}'],
            ["教材版本", packet.get("textbook", "—")],
            ["单元", packet.get("chapter", "—")],
            ["课文", packet.get("lesson", "—")],
            ["课时", f'{packet.get("lesson_hours", 1)}课时（{packet.get("planned_minutes", 45)}分钟）'],
            ["授课对象", packet.get("grade", "—")],
        ],
        widths=[3.5, 11.0],
    )

    # 二、教材与学情分析
    d.h1("二、教材与学情分析")
    d.h2("（一）教材分析")
    for para in nv.get("textbook_analysis", []):
        d.body(para)
    d.h2("（二）学情分析")
    for para in nv.get("learner_analysis", []):
        d.body(para)

    # 三、教学目标
    d.h1("三、教学目标")
    for obj in packet.get("objectives", []):
        d.body(obj)

    # 四、教学重点与难点
    d.h1("四、教学重点与难点")
    d.h2("（一）教学重点")
    for k in packet.get("key_points", []):
        d.body(k)
    d.h2("（二）教学难点")
    for k in packet.get("difficult_points", []):
        d.body(k)

    # 五、教学方法与准备
    d.h1("五、教学方法与教学准备")
    if nv.get("teaching_methods"):
        d.body(f'教学方法：{nv["teaching_methods"]}')
    if nv.get("preparation"):
        d.body(f'教学准备：{nv["preparation"]}')

    # 六、教学过程
    d.h1("六、教学过程")
    steps = packet.get("teaching_steps", [])
    rows = [
        [
            s.get("step"),
            s.get("name"),
            f'{s.get("minutes")}分钟',
            s.get("teacher_action", ""),
            s.get("student_action", ""),
            s.get("assessment", ""),
        ]
        for s in steps
    ]
    d.table(["环节", "教学环节", "时间", "教师活动", "学生活动", "评价要点"], rows,
            widths=[1.2, 2.6, 1.3, 4.4, 3.0, 2.0])
    total = sum(int(s.get("minutes", 0)) for s in steps)
    d.body(f"合计 {total} 分钟。", indent=False)

    # 七、板书设计
    if nv.get("blackboard_design"):
        d.h1("七、板书设计")
        for line in nv["blackboard_design"].split("\n"):
            d.body(line, indent=False)

    # 八、作业布置
    if nv.get("homework"):
        d.h1("八、作业布置")
        d.body(nv["homework"])

    # 九、AI 赋能设计
    ai = packet.get("ai_plan") or packet.get("ai_tasks") or []
    if ai:
        d.h1("九、AI 赋能设计")
        for item in ai:
            d.body(f'{item.get("id")}（{item.get("type")}）：{item.get("title")}', indent=False)
            d.body(f'课堂作用：{item.get("classroom_role", "")}')
            d.body(f'备用方案：{item.get("fallback", "")}')
            d.body(f'课堂核验：{item.get("verify", "")}')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lesson-packet", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    packet = json.loads(Path(args.lesson_packet).read_text(encoding="utf-8"))
    d = Doc()
    build(packet, d)
    d.save(args.out)
    print(f"[ok] 教学设计已生成: {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
